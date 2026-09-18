from contextlib import contextmanager
from pathlib import Path
from io import BytesIO
from datetime import datetime
import json, os, sqlite3, uuid, zipfile
from fastapi import FastAPI, UploadFile, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field
from typing import List
from openpyxl import load_workbook, Workbook
from .engine import validate, analyze, calculate, demo_book

ROOT=Path(os.environ.get('PROFIT_DATA_DIR','.data')); ROOT.mkdir(exist_ok=True,parents=True)
app=FastAPI(title='Profit Insight AI',version='0.1.0')

@contextmanager
def db():
    c=sqlite3.connect(ROOT/'profit.db')
    c.execute('CREATE TABLE IF NOT EXISTS records (id TEXT PRIMARY KEY, kind TEXT, payload TEXT, created_at TEXT)')
    try:
        with c:
            yield c
    finally:
        c.close()

def save(kind,payload,id=None):
    id=id or str(uuid.uuid4())
    with db() as c: c.execute('INSERT OR REPLACE INTO records VALUES (?,?,?,?)',(id,kind,json.dumps(payload,ensure_ascii=False,default=str),datetime.now().isoformat()))
    return id

def get(id,kind):
    with db() as c: row=c.execute('SELECT payload FROM records WHERE id=? AND kind=?',(id,kind)).fetchone()
    if not row: raise HTTPException(404,'대상을 찾을 수 없습니다.')
    return json.loads(row[0])

def workbook(file_id): return demo_book() if file_id=='demo' else get(file_id,'file')['book']

class Selection(BaseModel):
    file_id: str='demo'
    sales_sheet: str='매출_정상'
    cost_sheet: str='원가_정상'
    period: str=''
    region: str=''
    family: str=''

class Simulation(Selection):
    product_code: str
    customer_id: str
    quantity: float=Field(gt=0,allow_inf_nan=False)
    discount_rates: List[float]=[0,.05,.1]
    cost_month: str
    shipping_cost: float=Field(default=0,ge=0,allow_inf_nan=False)
    custom_cost: float=Field(default=0,ge=0,allow_inf_nan=False)
    cost_change: float=Field(default=0,gt=-1,le=10,allow_inf_nan=False)

class ReportRequest(BaseModel):
    analysis_id: str

@app.get('/api/health')
def health(): return {'status':'ok','mode':'local-single-user'}

@app.get('/api/files')
def files():
    with db() as c: rows=c.execute("SELECT id,payload,created_at FROM records WHERE kind='file' ORDER BY created_at DESC").fetchall()
    return [dict(file_id=id,file_name=json.loads(p)['file_name'],created_at=t,sheets=list(json.loads(p)['book']),counts={s:len(rs) for s,rs in json.loads(p)['book'].items()}) for id,p,t in rows]

@app.post('/api/files/upload')
async def upload(file: UploadFile):
    if not (file.filename or '').lower().endswith('.xlsx'): raise HTTPException(400,'.xlsx 파일만 지원합니다.')
    raw=await file.read(10*1024*1024+1)
    if len(raw)>10*1024*1024: raise HTTPException(413,'파일은 10MB 이하로 업로드하세요.')
    try:
        with zipfile.ZipFile(BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist())>100*1024*1024: raise ValueError('too large')
        wb=load_workbook(BytesIO(raw),data_only=True,read_only=True)
        book={}
        for ws in wb:
            if ws.max_row and ws.max_row>50001: raise ValueError('too many rows')
            if ws.max_column and ws.max_column>100: raise ValueError('too many columns')
            rows=ws.iter_rows(values_only=True); headers=next(rows,())
            names=[str(h).strip() if h is not None else '' for h in headers]
            if len([n for n in names if n])!=len(set(n for n in names if n)): raise ValueError('duplicate headers')
            book[ws.title]=[{k:(v.strftime('%Y-%m-%d') if isinstance(v,datetime) else v) for k,v in zip(names,row) if k} for row in rows if any(v is not None for v in row)]
        wb.close()
    except Exception: raise HTTPException(400,'Excel을 읽을 수 없습니다. 중복 헤더, 파일 손상 또는 허용 크기 초과 여부를 확인하세요.')
    id=str(uuid.uuid4()); (ROOT/f'{id}.xlsx').write_bytes(raw)
    save('file',dict(file_name=Path(file.filename).name,book=book),id)
    return dict(file_id=id,file_name=Path(file.filename).name,sheets=list(book),counts={s:len(r) for s,r in book.items()},status='uploaded')

@app.get('/api/files/template')
def template():
    wb=Workbook(); wb.remove(wb.active)
    for name,rows in demo_book().items():
        ws=wb.create_sheet(name); ws.append(list(rows[0]))
        for row in rows: ws.append(list(row.values()))
        ws.freeze_panes='A2'
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=22
    out=BytesIO(); wb.save(out)
    return Response(out.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="profit-insight-demo.xlsx"'})

@app.post('/api/validation/run')
def validation(req:Selection):
    result=validate(workbook(req.file_id),req.sales_sheet,req.cost_sheet)
    save('validation',dict(file_id=req.file_id,**result)); return result

@app.post('/api/profit/analyze')
def analysis(req:Selection):
    book=workbook(req.file_id)
    try: result=analyze(book,req.period,req.region,req.family,req.sales_sheet,req.cost_sheet)
    except ValueError as e: raise HTTPException(422,str(e))
    result.update(file_id=req.file_id,is_demo=req.file_id=='demo',products=book['제품마스터'],customers=book['고객마스터'],months=sorted({r['기준월'] for r in book[req.cost_sheet]}),costs=book[req.cost_sheet],filters=dict(region=req.region,family=req.family))
    result['analysis_id']=save('analysis',result)
    return result

@app.post('/api/simulation/discount')
def simulation(req:Simulation):
    import math
    if not 1<=len(req.discount_rates)<=20 or any(not math.isfinite(d) or not 0<=d<1 for d in req.discount_rates): raise HTTPException(422,'할인율은 0 이상 1 미만, 최대 20개여야 합니다.')
    book=workbook(req.file_id)
    if validate(book,req.sales_sheet,req.cost_sheet)['error_count']: raise HTTPException(422,'검증된 데이터만 시뮬레이션할 수 있습니다.')
    p=next((r for r in book['제품마스터'] if r['제품코드']==req.product_code),None)
    c=next((r for r in book[req.cost_sheet] if r['제품코드']==req.product_code and r['기준월']==req.cost_month),None)
    if not p or not c or not any(r['고객ID']==req.customer_id for r in book['고객마스터']): raise HTTPException(422,'제품·고객 또는 해당 월 원가가 없습니다.')
    from decimal import Decimal
    cost=Decimal(str(c['표준단위원가_KRW']))*(1+Decimal(str(req.cost_change)))
    return dict(unit_cost=float(cost),base_price=p['기본판매단가_KRW'],results=[dict(discount_rate=d,**calculate(req.quantity,p['기본판매단가_KRW'],d,cost,req.shipping_cost,req.custom_cost)) for d in req.discount_rates])

@app.post('/api/ai/report')
def report(req:ReportRequest):
    a=get(req.analysis_id,'analysis')
    if not a['count']: raise HTTPException(422,'선택한 조건에 분석 데이터가 없습니다.')
    top=max(a['by_product'],key=lambda r:r['adjusted_margin']); low=min(a['by_customer'],key=lambda r:r['adjusted_margin'])
    result=dict(title=f"{a['period']} 제품·고객별 수익성 분석 보고서",mode='rule_based',analysis_id=req.analysis_id,is_demo=a['is_demo'],executive_summary=f"총 매출 {a['revenue']:,.0f}원, 조정이익 {a['adjusted_profit']:,.0f}원, 조정이익률 {a['adjusted_margin']:.2%}입니다. 검증된 거래 {a['count']}건을 기준으로 집계했습니다.",key_findings=[f"{r['product_name']}({r['id']}): 매출 {r['revenue']:,.0f}원, 조정이익 {r['adjusted_profit']:,.0f}원, 조정이익률 {r['adjusted_margin']:.2%}." for r in a['by_product']],risk_points=[f"{low['customer_name']}의 조정이익률은 {low['adjusted_margin']:.2%}, 거래별 단순 평균 할인율은 {low['avg_discount']:.2%}입니다. 배송비 {low['shipping']:,.0f}원과 사양변경비 {low['custom']:,.0f}원이 차감되었습니다."],recommended_actions=[f"{low['customer_name']} 추가 할인 전 시뮬레이션으로 목표 이익률을 확인하세요.",f"{top['product_name']}의 제품 믹스 확대 가능성을 검토하세요.",'조정이익은 공통 판관비·연구개발비·이자·법인세를 제외한 관리회계 지표입니다. 최종 판단은 CFO 검토가 필요합니다.'],generated_at=datetime.now().isoformat())
    result['report_id']=save('report',result); return result
